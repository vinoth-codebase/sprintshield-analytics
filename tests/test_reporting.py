"""Tests for the Excel/PPTX exporters and the dashboard view layer."""

from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from sprintshield.analytics.monte_carlo import simulate_all_towers
from sprintshield.dashboard.views import (
    build_active_wip_table,
    build_capacity_donut,
    build_confidence_curve,
    build_support_tax_bar,
    build_tower_radar,
    build_velocity_trend,
    load_dashboard_data,
    simulate_what_if,
)
from sprintshield.reporting.excel_exporter import (
    ExcelExportResult,
    SprintSupportLedgerExporter,
    export_support_ledger,
)
from sprintshield.reporting.pptx_exporter import (
    PptxExportResult,
    SteerCoDeckExporter,
    export_steerco_deck,
)


# --------------------------------------------------------------------------- #
# Excel exporter
# --------------------------------------------------------------------------- #
@pytest.fixture()
def ledger(config, snapshots, tickets, tmp_path) -> ExcelExportResult:
    return export_support_ledger(
        config, snapshots, tickets, output_path=tmp_path / "ledger.xlsx"
    )


def test_ledger_creates_expected_sheets(ledger):
    assert ledger.path.exists()
    assert ledger.sheets == [
        "Executive KPI Summary",
        "Tower Breakdown Matrix",
        "Support Tax Heatmap",
        "Sprint Detail",
    ]
    assert ledger.rows_written > 0
    assert ledger.generated_at


def test_ledger_result_as_dict(ledger):
    payload = ledger.as_dict()
    assert payload["sheets"] == ledger.sheets
    assert payload["rows_written"] == ledger.rows_written


def test_ledger_workbook_is_readable_and_styled(ledger):
    wb = load_workbook(ledger.path)
    ws = wb["Executive KPI Summary"]
    assert ws["A1"].value.startswith("SprintShield")
    # A native bar chart should have been embedded.
    assert len(ws._charts) >= 1


def test_ledger_heatmap_has_conditional_formatting(ledger):
    wb = load_workbook(ledger.path)
    ws = wb["Support Tax Heatmap"]
    assert len(ws.conditional_formatting._cf_rules) >= 1


def test_ledger_tower_matrix_has_asset_breakdown(ledger):
    wb = load_workbook(ledger.path)
    ws = wb["Tower Breakdown Matrix"]
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("Support Burden by Business Asset" in v for v in values)


def test_ledger_default_output_path(config, snapshots, tickets, tmp_path, monkeypatch):
    monkeypatch.setenv("REPORT_OUTPUT_DIR", str(tmp_path / "reports"))
    from sprintshield.config import load_config

    cfg = load_config(project_root=config.project_root)
    result = SprintSupportLedgerExporter(cfg).export(snapshots, tickets)
    assert result.path.exists()
    assert result.path.name == "sprint_support_ledger.xlsx"


def test_ledger_handles_empty_frames(config, tmp_path):
    result = export_support_ledger(
        config, pd.DataFrame(), pd.DataFrame(), output_path=tmp_path / "empty.xlsx"
    )
    assert result.path.exists()
    wb = load_workbook(result.path)
    assert "Executive KPI Summary" in wb.sheetnames


def test_ledger_custom_tower_names(config, snapshots, tickets, tmp_path):
    result = export_support_ledger(
        config,
        snapshots,
        tickets,
        output_path=tmp_path / "named.xlsx",
        tower_names={"DE": "Platform Engineering"},
    )
    wb = load_workbook(result.path)
    ws = wb["Executive KPI Summary"]
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Platform Engineering" in values


# --------------------------------------------------------------------------- #
# PPTX exporter
# --------------------------------------------------------------------------- #
@pytest.fixture()
def deck(config, snapshots, tickets, tmp_path) -> PptxExportResult:
    _, summary = simulate_all_towers(config, snapshots, tickets, runs=300)
    return export_steerco_deck(
        config, snapshots, tickets, simulation_summary=summary, output_path=tmp_path / "deck.pptx"
    )


def test_deck_creates_three_slides(deck):
    assert deck.path.exists()
    assert deck.slides == 3
    assert deck.generated_at


def test_deck_result_as_dict(deck):
    assert deck.as_dict()["slides"] == 3


def test_deck_slide_titles(deck):
    prs = Presentation(deck.path)
    titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                titles.append(shape.text_frame.text)
    joined = "\n".join(titles)
    assert "Executive KPI Scorecard" in joined
    assert "Root-Cause Waterfall" in joined
    assert "Monte Carlo Capacity Recommendation" in joined


def test_deck_embeds_charts(deck):
    prs = Presentation(deck.path)
    chart_count = sum(
        1 for slide in prs.slides for shape in slide.shapes if shape.has_chart
    )
    assert chart_count >= 3


def test_deck_embeds_tables(deck):
    prs = Presentation(deck.path)
    table_count = sum(1 for slide in prs.slides for shape in slide.shapes if shape.has_table)
    assert table_count >= 2


def test_deck_without_supplied_simulation_runs_its_own(config, snapshots, tickets, tmp_path):
    result = SteerCoDeckExporter(config).export(
        snapshots, tickets, output_path=tmp_path / "auto.pptx"
    )
    assert result.slides == 3


def test_deck_default_output_path(config, snapshots, tickets, tmp_path, monkeypatch):
    monkeypatch.setenv("REPORT_OUTPUT_DIR", str(tmp_path / "reports"))
    from sprintshield.config import load_config

    cfg = load_config(project_root=config.project_root)
    result = SteerCoDeckExporter(cfg).export(snapshots, tickets)
    assert result.path.name == "stakeholder_steerco_deck.pptx"
    assert result.path.exists()


def test_deck_handles_empty_frames(config, tmp_path):
    result = export_steerco_deck(
        config, pd.DataFrame(), pd.DataFrame(), output_path=tmp_path / "empty.pptx"
    )
    assert result.path.exists()


# --------------------------------------------------------------------------- #
# Dashboard view layer
# --------------------------------------------------------------------------- #
def test_load_dashboard_data(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    assert not data.snapshots.empty
    assert not data.tickets.empty
    assert len(data.simulations) == len(config.towers)
    assert data.tower_names()["DE"] == "Data Engineering"
    assert "overall_support_tax_pct" in data.summary


def test_confidence_curve_figure(config, snapshots, tickets):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    figure = build_confidence_curve(data.simulations["DE"], planned_hours=200.0)
    assert figure.data
    assert len(figure.data) >= 2


def test_support_tax_bar_figure(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    figure = build_support_tax_bar(data.tower_metrics, data.tower_names())
    assert figure.data
    assert len(figure.data[0].x) == 4


def test_tower_radar_figure(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    figure = build_tower_radar(data.tower_metrics)
    assert len(figure.data) == 4


def test_capacity_donut_figure(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    figure = build_capacity_donut(data.tower_metrics, "DE")
    assert figure.data
    assert figure.data[0].hole == pytest.approx(0.62)


def test_capacity_donut_with_unknown_tower_falls_back(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    figure = build_capacity_donut(data.tower_metrics, "ZZ")
    assert figure.data


def test_velocity_trend_figure(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    assert build_velocity_trend(data.snapshots).data
    assert build_velocity_trend(data.snapshots, "DE").data


def test_view_builders_handle_empty_frames():
    assert build_support_tax_bar(pd.DataFrame()).data == ()
    assert build_tower_radar(pd.DataFrame()).data == ()
    assert build_capacity_donut(pd.DataFrame()).data == ()
    assert build_velocity_trend(pd.DataFrame()).data == ()


def test_active_wip_table(tickets):
    wip = build_active_wip_table(tickets)
    assert not wip.empty
    assert "priority" in wip.columns
    assert {"ticket_id", "tower_id", "asset", "priority", "summary"} <= set(wip.columns)
    # Rows are ordered by severity: the priority rank must be non-decreasing.
    rank = {"P1": 0, "P2": 1, "P3": 2, "P4": 3, "P5": 4}
    ranks = wip["priority"].map(rank).tolist()
    assert ranks == sorted(ranks)
    assert wip["priority"].iloc[0] in {"P1", "P2", "P3"}


def test_active_wip_table_empty():
    assert build_active_wip_table(pd.DataFrame()).empty
    resolved = pd.DataFrame({"status": ["Done"], "ticket_id": ["x"], "tower_id": ["DE"], "asset": ["a"], "priority": ["P1"], "summary": ["s"], "sprint_id": ["S1"]})
    assert build_active_wip_table(resolved).empty


def test_simulate_what_if(config):
    data = load_dashboard_data(config=config, sprints=6, runs=300)
    result = simulate_what_if(data, "DE", planned_points=60)
    assert result.runs == config.simulation.runs
    assert 0.0 <= result.delivery_confidence() <= 1.0


# --------------------------------------------------------------------------- #
# Streamlit cockpit smoke test
# --------------------------------------------------------------------------- #
def test_streamlit_app_renders_without_exceptions(project_root, monkeypatch):
    """Drive the real Streamlit app through Streamlit's own test harness."""
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("DELTA_LAKE_PATH", str(project_root / "data" / "delta"))

    app = AppTest.from_file(
        str(project_root / "src" / "sprintshield" / "dashboard" / "app.py"),
        default_timeout=240,
    )
    app.run()

    assert not app.exception, [str(e.value) for e in app.exception]
    assert len(app.tabs) == 4
    assert len(app.metric) >= 5
    assert [s.label for s in app.selectbox] == ["Tower"]
    assert any("Planned sprint commitment" in s.label for s in app.slider)
    assert any("Build Excel support ledger" in b.label for b in app.button)

    # Moving the What-If slider must re-simulate and update delivery confidence.
    planned = next(s for s in app.slider if "Planned sprint commitment" in s.label)
    planned.set_value(40).run()
    assert not app.exception, [str(e.value) for e in app.exception]
    confidence = next(m for m in app.metric if m.label == "Delivery confidence")
    assert float(confidence.value.rstrip("%")) > 80.0
